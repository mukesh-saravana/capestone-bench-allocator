from __future__ import annotations

import csv
from contextlib import asynccontextmanager
from datetime import date, datetime
from io import BytesIO, StringIO
from uuid import uuid4

from fastapi import Cookie, Depends, FastAPI, File, Header, HTTPException, Response, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook

from .data import store
from .models import (
    AllocationsSummary,
    AllocationHistory,
    AssignRequest,
    BenchMetrics,
    BenchTrendPoint,
    CandidateImportResult,
    ChatQueryRequest,
    ChatQueryResponse,
    DataSetSummary,
    Employee,
    LoginRequest,
    LoginResponse,
    ProjectNeed,
    RagReindexResponse,
    RagStatusResponse,
    RecommendationRequest,
    RecommendationResponse,
    SkillTag,
    SkillTagCreateRequest,
    SkillTagUpdateRequest,
    User,
    UtilizationDepartmentPoint,
    UtilizationMetrics,
    UtilizationTrendPoint,
)
from .rag import RAGService
from .settings import get_settings
from .services import recommend, recommend_from_query

settings = get_settings()
rag_service = RAGService(store=store, settings=settings)


@asynccontextmanager
async def lifespan(_: FastAPI):
    store.initialize()
    rag_service.refresh_index()
    yield


app = FastAPI(title="Capstone Bench Allocator Backend", version="0.1.0", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[settings.cors_origin],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def _extract_bearer_token(authorization: str | None) -> str | None:
    if not authorization:
        return None
    prefix = "bearer "
    if authorization.lower().startswith(prefix):
        return authorization[len(prefix) :].strip()
    return None


def get_current_user(
    authorization: str | None = Header(default=None),
    auth_token: str | None = Cookie(default=None, alias=settings.auth_cookie_name),
) -> User:
    token = _extract_bearer_token(authorization) or auth_token
    user = store.get_user_by_token(token)
    if not user:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    return user


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/api/rag/status", response_model=RagStatusResponse)
def rag_status(current_user: User = Depends(get_current_user)) -> RagStatusResponse:
    _ = current_user
    return RagStatusResponse(**rag_service.status())


@app.post("/api/rag/reindex", response_model=RagReindexResponse)
def rag_reindex(current_user: User = Depends(get_current_user)) -> RagReindexResponse:
    _ = current_user
    rag_service.refresh_index()
    state = rag_service.status()
    return RagReindexResponse(status="ok", indexedChunks=int(state["indexedChunks"]), retrievalMode=str(state["retrievalMode"]))


@app.post("/api/auth/login", response_model=LoginResponse)
def login(payload: LoginRequest, response: Response) -> LoginResponse:
    if len(payload.password) < 4:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Password must be at least 4 characters")

    user = store.get_user_by_email(payload.email)
    if not user:
        user = User(id=f"usr-{uuid4()}", name=payload.email.split("@")[0].title(), email=payload.email, role="resource_manager")
        user = store.save_user(user)

    token = store.issue_token(payload.email)
    response.set_cookie(settings.auth_cookie_name, token, httponly=True, samesite="lax")
    return LoginResponse(token=token, user=user)


@app.post("/api/auth/logout")
def logout(
    response: Response, authorization: str | None = Header(default=None), auth_token: str | None = Cookie(default=None, alias=settings.auth_cookie_name)
) -> dict[str, str]:
    token = _extract_bearer_token(authorization) or auth_token
    if token:
        store.revoke_token(token)
    response.delete_cookie(settings.auth_cookie_name)
    return {"status": "ok"}


@app.get("/api/auth/me", response_model=User)
def me(current_user: User = Depends(get_current_user)) -> User:
    return current_user


@app.get("/api/settings/skills", response_model=list[SkillTag])
def list_skill_tags(current_user: User = Depends(get_current_user)) -> list[SkillTag]:
    _ = current_user
    return store.list_skill_tags()


@app.post("/api/settings/skills", response_model=SkillTag, status_code=status.HTTP_201_CREATED)
def create_skill_tag(payload: SkillTagCreateRequest, current_user: User = Depends(get_current_user)) -> SkillTag:
    _ = current_user
    try:
        return store.create_skill_tag(name=payload.name, category=payload.category)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@app.put("/api/settings/skills/{skill_id}", response_model=SkillTag)
def update_skill_tag(skill_id: str, payload: SkillTagUpdateRequest, current_user: User = Depends(get_current_user)) -> SkillTag:
    _ = current_user
    try:
        return store.update_skill_tag(skill_id=skill_id, name=payload.name, category=payload.category)
    except ValueError as exc:
        detail = str(exc)
        status_code = status.HTTP_404_NOT_FOUND if "not found" in detail.lower() else status.HTTP_400_BAD_REQUEST
        raise HTTPException(status_code=status_code, detail=detail) from exc


@app.delete("/api/settings/skills/{skill_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_skill_tag(skill_id: str, current_user: User = Depends(get_current_user)) -> Response:
    _ = current_user
    try:
        store.delete_skill_tag(skill_id=skill_id)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc)) from exc
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@app.get("/api/employees", response_model=list[Employee])
def list_employees(current_user: User = Depends(get_current_user)) -> list[Employee]:
    _ = current_user
    return store.list_employees()


@app.get("/api/project-needs", response_model=list[ProjectNeed])
def list_project_needs(current_user: User = Depends(get_current_user)) -> list[ProjectNeed]:
    _ = current_user
    return store.list_project_needs()


@app.get("/api/allocation-history", response_model=list[AllocationHistory])
def list_allocation_history(current_user: User = Depends(get_current_user)) -> list[AllocationHistory]:
    _ = current_user
    return store.list_allocations()


@app.get("/api/data/summary", response_model=list[DataSetSummary])
def dataset_summary(current_user: User = Depends(get_current_user)) -> list[DataSetSummary]:
    _ = current_user
    return [DataSetSummary(**item) for item in store.get_dataset_summaries()]


def _read_tabular_rows(file_name: str, content: bytes) -> list[dict[str, object]]:
    normalized_name = file_name.lower()
    if normalized_name.endswith(".csv"):
        text_data = content.decode("utf-8-sig")
        reader = csv.DictReader(StringIO(text_data))
        rows = list(reader)
    elif normalized_name.endswith(".xlsx"):
        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value).strip() if value is not None else "" for value in values[0]]
        rows = []
        for value_row in values[1:]:
            row = {headers[idx]: value_row[idx] for idx in range(min(len(headers), len(value_row)))}
            rows.append(row)
    else:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only CSV and XLSX files are supported")
    return rows


def _parse_optional_date(value: object, *, field: str, row_number: int) -> date | None:
    if value in (None, ""):
        return None
    try:
        if isinstance(value, date):
            if isinstance(value, datetime):
                return value.date()
            return value
        return date.fromisoformat(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=f"Invalid date for '{field}' at row {row_number}; expected YYYY-MM-DD"
        ) from exc


def _parse_candidate_rows(file_name: str, content: bytes) -> list[dict[str, object]]:
    rows = _read_tabular_rows(file_name, content)

    parsed: list[dict[str, object]] = []
    required_columns = {
        "candidate_id",
        "name",
        "email",
        "role",
        "department",
        "experience_years",
        "skills",
        "availability",
        "utilization_pct",
        "interview_score",
        "interview_result",
    }
    allowed_availability = {"available", "allocated", "on_leave", "exiting"}
    for index, row in enumerate(rows, start=2):
        normalized = {str(key).strip().lower(): row.get(key) for key in row.keys() if key is not None}
        missing = [column for column in required_columns if column not in normalized]
        if missing:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Missing columns: {', '.join(sorted(missing))}")
        try:
            availability = str(normalized["availability"]).strip()
            if availability not in allowed_availability:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Invalid availability '{availability}' at row {index}",
                )
            parsed.append(
                {
                    "id": str(normalized["candidate_id"]).strip(),
                    "name": str(normalized["name"]).strip(),
                    "email": str(normalized["email"]).strip(),
                    "role": str(normalized["role"]).strip(),
                    "department": str(normalized["department"]).strip(),
                    "experienceYears": int(normalized["experience_years"]),
                    "skills": str(normalized["skills"]).strip(),
                    "availability": availability,
                    "utilizationPct": int(normalized["utilization_pct"]),
                    "interviewScore": int(normalized["interview_score"]) if normalized["interview_score"] not in (None, "") else None,
                    "interviewResult": str(normalized["interview_result"]).strip() if normalized["interview_result"] not in (None, "") else None,
                    "benchSince": None,
                }
            )
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Invalid data at row {index}") from exc
    return parsed


def _parse_employee_rows(file_name: str, content: bytes) -> list[dict[str, object]]:
    rows = _read_tabular_rows(file_name, content)
    parsed: list[dict[str, object]] = []
    required_columns = {"employee_id", "name", "email", "role", "department", "experience_years", "skills", "availability", "utilization_pct"}
    allowed_availability = {"available", "allocated", "on_leave", "exiting"}
    for index, row in enumerate(rows, start=2):
        normalized = {str(key).strip().lower(): row.get(key) for key in row.keys() if key is not None}
        missing = [column for column in required_columns if column not in normalized]
        if missing:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Missing columns: {', '.join(sorted(missing))}")
        availability = str(normalized["availability"]).strip()
        if availability not in allowed_availability:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid availability '{availability}' at row {index}",
            )
        try:
            parsed.append(
                {
                    "id": str(normalized["employee_id"]).strip(),
                    "name": str(normalized["name"]).strip(),
                    "email": str(normalized["email"]).strip(),
                    "role": str(normalized["role"]).strip(),
                    "department": str(normalized["department"]).strip(),
                    "experienceYears": int(normalized["experience_years"]),
                    "skills": str(normalized["skills"]).strip(),
                    "availability": availability,
                    "utilizationPct": int(normalized["utilization_pct"]),
                    "benchSince": _parse_optional_date(normalized.get("bench_since"), field="bench_since", row_number=index),
                }
            )
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Invalid data at row {index}") from exc
    return parsed


def _parse_project_need_rows(file_name: str, content: bytes) -> list[dict[str, object]]:
    rows = _read_tabular_rows(file_name, content)
    parsed: list[dict[str, object]] = []
    required_columns = {"need_id", "project_id", "project_name", "role_title", "open_slots", "required_skills", "status", "priority"}
    allowed_status = {"open", "filled", "cancelled"}
    allowed_priority = {"high", "medium", "low"}
    for index, row in enumerate(rows, start=2):
        normalized = {str(key).strip().lower(): row.get(key) for key in row.keys() if key is not None}
        missing = [column for column in required_columns if column not in normalized]
        if missing:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Missing columns: {', '.join(sorted(missing))}")
        status_value = str(normalized["status"]).strip()
        priority = str(normalized["priority"]).strip()
        if status_value not in allowed_status:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Invalid status '{status_value}' at row {index}")
        if priority not in allowed_priority:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Invalid priority '{priority}' at row {index}")
        try:
            parsed.append(
                {
                    "id": str(normalized["need_id"]).strip(),
                    "projectId": str(normalized["project_id"]).strip(),
                    "projectName": str(normalized["project_name"]).strip(),
                    "roleTitle": str(normalized["role_title"]).strip(),
                    "openSlots": int(normalized["open_slots"]),
                    "requiredSkills": str(normalized["required_skills"]).strip(),
                    "startDate": _parse_optional_date(normalized.get("start_date"), field="start_date", row_number=index),
                    "status": status_value,
                    "priority": priority,
                }
            )
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Invalid data at row {index}") from exc
    return parsed


def _parse_allocation_rows(file_name: str, content: bytes) -> list[dict[str, object]]:
    rows = _read_tabular_rows(file_name, content)
    parsed: list[dict[str, object]] = []
    required_columns = {
        "allocation_id",
        "employee_id",
        "employee_name",
        "project_id",
        "project_name",
        "role",
        "start_date",
        "outcome",
    }
    allowed_outcome = {"completed", "transferred", "early_exit", "ongoing"}
    for index, row in enumerate(rows, start=2):
        normalized = {str(key).strip().lower(): row.get(key) for key in row.keys() if key is not None}
        missing = [column for column in required_columns if column not in normalized]
        if missing:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Missing columns: {', '.join(sorted(missing))}")
        outcome = str(normalized["outcome"]).strip()
        if outcome not in allowed_outcome:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Invalid outcome '{outcome}' at row {index}")
        parsed.append(
            {
                "id": str(normalized["allocation_id"]).strip(),
                "employeeId": str(normalized["employee_id"]).strip(),
                "employeeName": str(normalized["employee_name"]).strip(),
                "projectId": str(normalized["project_id"]).strip(),
                "projectName": str(normalized["project_name"]).strip(),
                "role": str(normalized["role"]).strip(),
                "startDate": _parse_optional_date(normalized["start_date"], field="start_date", row_number=index),
                "endDate": _parse_optional_date(normalized.get("end_date"), field="end_date", row_number=index),
                "outcome": outcome,
                "notes": str(normalized["notes"]).strip() if normalized.get("notes") not in (None, "") else None,
            }
        )
        if parsed[-1]["startDate"] is None:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Missing start_date at row {index}")
    return parsed


@app.post("/api/import/{dataset_key}", response_model=CandidateImportResult)
async def import_dataset(dataset_key: str, file: UploadFile = File(...), current_user: User = Depends(get_current_user)) -> CandidateImportResult:
    _ = current_user
    file_name = file.filename or f"{dataset_key}.csv"
    content = await file.read()

    if dataset_key == "candidate_profiles":
        rows = _parse_candidate_rows(file_name, content)
        result = store.import_candidates(rows, file_name)
    elif dataset_key == "employees":
        rows = _parse_employee_rows(file_name, content)
        result = store.import_employees(rows, file_name)
    elif dataset_key == "project_needs":
        rows = _parse_project_need_rows(file_name, content)
        result = store.import_project_needs(rows, file_name)
    elif dataset_key == "allocation_history":
        rows = _parse_allocation_rows(file_name, content)
        result = store.import_allocation_history(rows, file_name)
    else:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Unsupported dataset '{dataset_key}'")
    rag_service.refresh_index()
    return CandidateImportResult(**result)


@app.post("/api/import/candidates", response_model=CandidateImportResult)
async def import_candidates(file: UploadFile = File(...), current_user: User = Depends(get_current_user)) -> CandidateImportResult:
    return await import_dataset(dataset_key="candidate_profiles", file=file, current_user=current_user)


@app.get("/api/export/candidates")
def export_candidates(current_user: User = Depends(get_current_user)) -> StreamingResponse:
    _ = current_user
    payload = store.export_candidates_csv()
    return StreamingResponse(
        iter([payload.encode("utf-8")]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=candidates-export.csv"},
    )


@app.post("/api/chat/query", response_model=ChatQueryResponse)
def chat_query(payload: ChatQueryRequest, current_user: User = Depends(get_current_user)) -> ChatQueryResponse:
    try:
        rag_context = rag_service.retrieve(
            query=payload.query,
            department=(payload.filters.department if payload.filters else None),
            top_k=3,
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"RAG retrieval failed: {exc}") from exc

    recommendation_response = recommend_from_query(
        query=payload.query,
        strategy=payload.strategy,
        department=(payload.filters.department if payload.filters else None),
        evidence_snippets=rag_context.snippets,
    )
    recommendations = recommendation_response.recommendations
    top_k = len(recommendations)
    top_names = ", ".join(rec.employee.name for rec in recommendations[:3]) if recommendations else "No matching candidates"

    if top_k == 1:
        answer = f"For '{payload.query}', the best candidate is {top_names}. Retrieval mode: {rag_context.mode}."
    else:
        answer = f"For '{payload.query}', top candidates are: {top_names}. Retrieval mode: {rag_context.mode}."
    evidence: list[str] = []
    seen_snippets: set[str] = set()
    for snippet in [*rag_context.snippets, *(snippet for rec in recommendations for snippet in rec.evidenceSnippets)]:
        if snippet in seen_snippets:
            continue
        seen_snippets.add(snippet)
        evidence.append(snippet)
        if len(evidence) >= 5:
            break
    _ = current_user
    return ChatQueryResponse(
        answer=answer,
        recommendations=recommendations,
        evidenceSnippets=evidence,
        citations=rag_context.citations,
        retrievalMode=rag_context.mode,
        messageId=f"msg-{uuid4()}",
    )


@app.post("/api/recommendations", response_model=RecommendationResponse)
def recommendations(payload: RecommendationRequest, current_user: User = Depends(get_current_user)) -> RecommendationResponse:
    required_skills = payload.requiredSkills
    if payload.projectNeedId:
        need = store.project_needs.get(payload.projectNeedId)
        if not need:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Project need not found")
        required_skills = need.requiredSkills

    _ = current_user
    return recommend(required_skills=required_skills, department=payload.department, strategy=payload.strategy, top_k=5)


@app.get("/api/dashboard/bench", response_model=BenchMetrics)
def bench_metrics(current_user: User = Depends(get_current_user)) -> BenchMetrics:
    _ = current_user
    by_department = store.get_bench_department_counts()
    return BenchMetrics(
        totalOnBench=sum(by_department.values()),
        trend=[
            BenchTrendPoint(date=date.fromisoformat("2026-06-24"), count=8),
            BenchTrendPoint(date=date.fromisoformat("2026-07-01"), count=7),
            BenchTrendPoint(date=date.fromisoformat("2026-07-08"), count=6),
            BenchTrendPoint(date=date.fromisoformat("2026-07-15"), count=5),
            BenchTrendPoint(date=date.fromisoformat("2026-07-22"), count=4),
        ],
        byDepartment=[{"department": department, "count": count} for department, count in sorted(by_department.items())],
    )


@app.get("/api/dashboard/utilization", response_model=UtilizationMetrics)
def utilization_metrics(current_user: User = Depends(get_current_user)) -> UtilizationMetrics:
    _ = current_user
    employees = list(store.employees.values())
    average_pct = round(sum(emp.utilizationPct for emp in employees) / len(employees))

    per_department: dict[str, list[int]] = {}
    for emp in employees:
        per_department.setdefault(emp.department, []).append(emp.utilizationPct)

    by_department = [
        UtilizationDepartmentPoint(department=department, pct=round(sum(values) / len(values)))
        for department, values in sorted(per_department.items())
    ]
    return UtilizationMetrics(
        averagePct=average_pct,
        trend=[
            UtilizationTrendPoint(date=date.fromisoformat("2026-06-24"), pct=68),
            UtilizationTrendPoint(date=date.fromisoformat("2026-07-01"), pct=70),
            UtilizationTrendPoint(date=date.fromisoformat("2026-07-08"), pct=73),
            UtilizationTrendPoint(date=date.fromisoformat("2026-07-15"), pct=74),
            UtilizationTrendPoint(date=date.fromisoformat("2026-07-22"), pct=76),
        ],
        byDepartment=by_department,
    )


@app.get("/api/dashboard/allocations", response_model=AllocationsSummary)
def allocation_summary(current_user: User = Depends(get_current_user)) -> AllocationsSummary:
    _ = current_user
    total = len(store.allocations)
    active = sum(1 for allocation in store.allocations if allocation.outcome == "ongoing")
    completed = sum(1 for allocation in store.allocations if allocation.outcome == "completed")
    return AllocationsSummary(
        totalAllocations=total,
        activeAllocations=active,
        completedAllocations=completed,
        recent=store.allocations[:5],
    )


@app.post("/api/allocations")
def create_allocation(payload: AssignRequest, current_user: User = Depends(get_current_user)):
    _ = current_user
    if payload.employeeId not in store.employees:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Employee not found")
    if payload.projectNeedId not in store.project_needs:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Project need not found")
    try:
        allocation = store.create_allocation(
            employee_id=payload.employeeId,
            project_need_id=payload.projectNeedId,
            role=payload.role,
            start_date=payload.startDate,
            notes=payload.notes,
        )
        rag_service.refresh_index()
        return allocation
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc)) from exc
